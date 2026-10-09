"""VO_MANIFEST.md -> VO_MANIFEST.xlsx: the same manifest as a spreadsheet, for the team.

    python tools/vo_manifest_xlsx.py [VO_MANIFEST.md] [VO_MANIFEST.xlsx]

Three sheets:
  VO by screen  every row of the screens' tables, in play order; each screen's notes (the moments with no voiceover,
                where the game differs from the deck) as grey rows under it
  Clip list     each file in assets/voiceover once: its line, length, the screens that play it, its status
  Notes         the manifest's summary, the counting clips' timing, how the clips are made, the clips removed

Status comes from the Source column: Matches PPT, Re-record, Not in PPT (the transition cards), SFX.
VO_MANIFEST.md is the source - re-run this after any change to it. Needs openpyxl.
"""
import datetime, os, re, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "VO_MANIFEST.md")
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, "VO_MANIFEST.xlsx")
VO_DIR = os.path.join(os.path.dirname(os.path.abspath(SRC)), "assets", "voiceover")


def clean(s):
    """markdown -> plain text: bold, italics and code marks off"""
    s = s.replace("**", "").replace("`", "")
    return re.sub(r"\*([^*]+)\*", r"\1", s).strip()


# ---- read the manifest: its title, then sections of tables and text blocks (a paragraph or one list item) ----
title, sections = "", [{"head": None, "blocks": []}]
para = []

def flush():
    if para:
        sections[-1]["blocks"].append(("text", " ".join(x.strip() for x in para)))
        para.clear()

for ln in open(SRC, encoding="utf-8").read().splitlines():
    if ln.startswith("# "):
        title = clean(ln[2:])
    elif ln.startswith("## "):
        flush(); sections.append({"head": clean(ln[3:]), "blocks": []})
    elif ln.startswith("|"):
        flush()
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-+:?", c) for c in cells):
            continue                                                      # the |---|---| rule
        blocks = sections[-1]["blocks"]
        if blocks and blocks[-1][0] == "table":
            blocks[-1][1].append(cells)
        else:
            blocks.append(("table", [cells]))
    elif not ln.strip():
        flush()
    elif ln.startswith("- "):
        flush(); para.append(ln)                                          # a list item is a block of its own
    else:
        para.append(ln)
flush()


def split_head(h):
    """'T2 — measuring the table — PPT page 3 (slide 5)' -> ('T2 — measuring the table', 'Page 3 (slide 5)')"""
    m = re.match(r"(.*?) — PPT (pages? .*)$", h)
    return (m.group(1), m.group(2)[0].upper() + m.group(2)[1:]) if m else (h, "—")


def code(screen):
    """the screen's short name for the clip list: T2, G2b, 'P2, P3', Title, End, 'Tutorial card' ..."""
    m = re.match(r"(T\d+|G\d+b?|P\d+)(?: and (P\d+))?\b", screen)
    if m:
        return ", ".join(x for x in m.groups() if x)
    if screen.startswith("Title"):
        return "Title"
    if screen.startswith("End"):
        return "End"
    m = re.match(r"Transition into (?:the )?(\w+)", screen)
    return m.group(1).capitalize() + " card" if m else screen


def status(source):
    s = source.lower()
    if "re-record" in s:
        return "Re-record"
    if s.startswith("not in the ppt"):
        return "Not in PPT"
    if source.strip() in ("", "—", "-"):
        return "SFX"
    return "Matches PPT"


def clip_names(cell):
    return re.findall(r"`([^`]+)`", cell)


def clip_label(names):
    """vo_num_3 … vo_num_9 for a run of counting clips, else the names"""
    ks = [int(m.group(1)) for n in names for m in [re.fullmatch(r"vo_num_(\d+)", n)] if m]
    if len(names) >= 3 and len(ks) == len(names) and ks == list(range(ks[0], ks[0] + len(ks))):
        return "vo_num_%d … vo_num_%d" % (ks[0], ks[-1])
    return ", ".join(names)


def seconds(length):
    m = re.match(r"~?(\d+(?:\.\d+)?)s\b", length)
    return float(m.group(1)) if m else None


# ---- the rows ----
screen_rows, notes_rows, counting = [], [], {}
for sec in sections:
    tables = [b[1] for b in sec["blocks"] if b[0] == "table"]
    is_screen = any("When it plays" in t[0] for t in tables)
    if sec["head"] and not is_screen:
        for t in tables:                                                  # the counting clips' own table
            if t[0][:3] == ["Clip", "Line", "Length"]:
                for c in t[1:]:
                    for n in clip_names(c[0]):
                        counting[n] = {"line": clean(c[1]), "length": clean(c[2])}
    if is_screen:
        screen, page = split_head(sec["head"])
        for kind, b in sec["blocks"]:
            if kind == "table":
                for c in b[1:]:
                    names = clip_names(c[0])
                    screen_rows.append({"screen": screen, "page": page, "names": names, "clip": clip_label(names),
                                        "line": clean(c[1]), "when": clean(c[2]), "length": clean(c[3]),
                                        "source": clean(c[4]), "status": status(clean(c[4]))})
            else:
                note = clean(b)
                screen_rows.append({"screen": screen, "page": page, "names": [], "clip": "", "line": "",
                                    "when": note, "length": "", "source": "",
                                    "status": "Game ≠ deck" if "differs from the new deck" in note else "Note"})
    else:
        topic = sec["head"] or "About this manifest"
        for kind, b in sec["blocks"]:
            if kind == "text":
                t = clean(b)
                notes_rows.append((topic, "•  " + t[2:] if t.startswith("- ") else t))

clips, order = {}, []
for r in screen_rows:
    for n in r["names"]:
        if n not in clips:
            num = counting.get(n)
            clips[n] = {"line": num["line"] if num else r["line"], "length": num["length"] if num else r["length"],
                        "status": r["status"], "source": r["source"], "used": []}
            order.append(n)
        if code(r["screen"]) not in clips[n]["used"]:
            clips[n]["used"].append(code(r["screen"]))
nums = sorted((n for n in order if n in counting), key=lambda n: int(n.rsplit("_", 1)[1]))
order = [n for n in order if n not in counting] + nums                   # the counting clips last, एक … बारह

files = {}
if os.path.isdir(VO_DIR):
    for f in os.listdir(VO_DIR):
        files[os.path.splitext(f)[0]] = f
    missing, unused = sorted(set(order) - set(files)), sorted(set(files) - set(order))
    if missing:
        print("WARNING: in the manifest, no file:", ", ".join(missing))
    if unused:
        print("WARNING: files the manifest does not list:", ", ".join(unused))

# ---- the workbook ----
HEAD_FILL, HEAD_FONT = PatternFill("solid", fgColor="1F4E78"), Font(bold=True, color="FFFFFF")
FILLS = {"Matches PPT": "E2EFDA", "Re-record": "F8CBAD", "Not in PPT": "DDEBF7", "SFX": "EDEDED",
         "Note": "F2F2F2", "Game ≠ deck": "FCE4D6"}
HINDI = "Nirmala UI"                                                     # Windows' Devanagari UI font
TOP = Alignment(wrap_text=True, vertical="top")
RULE = Side(style="thin", color="8EA9DB")
stamp = datetime.date.fromtimestamp(os.path.getmtime(SRC)).isoformat()


def sheet(ws, heading, cols, rows, legend=True):
    """cols: [(name, width)]; rows: lists of cell values. Returns the first data row."""
    n = len(cols)
    ws.append([heading]); ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    ws["A1"].font = Font(bold=True, size=14); ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 24
    ws.append(["From VO_MANIFEST.md (%s)." % stamp + (" Status: Matches PPT · Re-record · Not in PPT (the transition "
               "cards) · SFX · Note (no voiceover / how the game behaves) · Game ≠ deck." if legend else "")])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    ws["A2"].font = Font(italic=True, color="595959"); ws["A2"].alignment = Alignment(vertical="top")
    ws.append([c[0] for c in cols])
    for i, (name, width) in enumerate(cols, 1):
        cell = ws.cell(row=3, column=i)
        cell.fill, cell.font, cell.alignment = HEAD_FILL, HEAD_FONT, Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A4"
    if rows and legend:
        ws.auto_filter.ref = "A3:%s%d" % (get_column_letter(n), 3 + len(rows))
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "3:3"
    for row in ws.iter_rows(min_row=4):
        for cell in row:
            cell.alignment = TOP
    return 4


wb = Workbook()

# 1. VO by screen
ws = wb.active; ws.title = "VO by screen"
cols = [("#", 5), ("Screen", 26), ("PPT page", 16), ("Clip", 22), ("Line", 46), ("When it plays", 54),
        ("Length", 10), ("Source", 34), ("Status", 13)]
rows, k, prev = [], 0, None
for r in screen_rows:
    vo = bool(r["names"])
    if vo:
        k += 1
    rows.append([k if vo else None, r["screen"], r["page"], r["clip"], r["line"], r["when"], r["length"],
                 r["source"], r["status"]])
first = sheet(ws, title, cols, rows)
for i, r in enumerate(screen_rows):
    row = first + i
    st = r["status"]
    ws.cell(row=row, column=5).font = Font(name=HINDI, size=12)
    ws.cell(row=row, column=9).fill = PatternFill("solid", fgColor=FILLS[st])
    if st == "Re-record":
        for c in range(1, 10):
            ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor="FCE4D6")
        ws.cell(row=row, column=9).fill = PatternFill("solid", fgColor=FILLS[st])
        ws.cell(row=row, column=9).font = Font(bold=True)
    elif not r["names"]:                                                  # a note row
        for c in range(1, 10):
            ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor=FILLS[st])
            ws.cell(row=row, column=c).font = Font(italic=True, color="404040", bold=(st == "Game ≠ deck" and c == 9))
    if r["screen"] != prev:                                               # a line between screens
        for c in range(1, 10):
            ws.cell(row=row, column=c).border = Border(top=RULE)
        prev = r["screen"]

# 2. Clip list
ws = wb.create_sheet("Clip list")
cols = [("#", 5), ("File", 26), ("Line", 46), ("Length (s)", 11), ("Used on", 30), ("Status", 13), ("Source", 52)]
rows = []
for i, n in enumerate(order, 1):
    c = clips[n]
    sec = seconds(c["length"])
    rows.append([i, files.get(n, n), c["line"], sec if sec is not None else c["length"], ", ".join(c["used"]),
                 c["status"], c["source"]])
first = sheet(ws, title + " — clip list (%d files)" % len(order), cols, rows)
for i, n in enumerate(order):
    row, st = first + i, clips[n]["status"]
    ws.cell(row=row, column=3).font = Font(name=HINDI, size=12)
    ws.cell(row=row, column=4).number_format = "0.0"
    ws.cell(row=row, column=4).alignment = Alignment(horizontal="center", vertical="top")
    ws.cell(row=row, column=6).fill = PatternFill("solid", fgColor=FILLS[st])
    if st == "Re-record":
        for c in range(1, 8):
            ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor="FCE4D6")
        ws.cell(row=row, column=6).fill = PatternFill("solid", fgColor=FILLS[st])
        ws.cell(row=row, column=6).font = Font(bold=True)

# 3. Notes
ws = wb.create_sheet("Notes")
rows, prev = [], None
for topic, note in notes_rows:
    rows.append([topic if topic != prev else None, note]); prev = topic
sheet(ws, title + " — notes", [("Topic", 26), ("Note", 120)], rows, legend=False)
for row in ws.iter_rows(min_row=4, max_col=1):
    for cell in row:
        cell.font = Font(bold=True)
        if cell.value:
            for c in range(1, 3):
                ws.cell(row=cell.row, column=c).border = Border(top=RULE)

wb.properties.title = title
wb.save(OUT)
by = {}
for n in order:
    by[clips[n]["status"]] = by.get(clips[n]["status"], 0) + 1
print("%s: %d rows by screen (%d VO), %d files (%s), %d notes" % (
    os.path.basename(OUT), len(screen_rows), sum(1 for r in screen_rows if r["names"]), len(order),
    ", ".join("%s %d" % kv for kv in sorted(by.items())), len(notes_rows)))
